"""
ReguAI Presentation Generator
Creates a professional PowerPoint presentation from the codebase analysis
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


def create_presentation():
    """Generate the complete ReguAI presentation"""
    prs = Presentation()
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)

    # Define color scheme
    PRIMARY_COLOR = RGBColor(37, 99, 235)  # Blue
    SECONDARY_COLOR = RGBColor(71, 85, 105)  # Slate
    ACCENT_COLOR = RGBColor(16, 185, 129)  # Green

    def add_title_slide(title, subtitle):
        """Add a title slide"""
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # Background
        background = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        background.fill.solid()
        background.fill.fore_color.rgb = PRIMARY_COLOR
        background.line.fill.background()

        # Title
        title_box = slide.shapes.add_textbox(
            Inches(1), Inches(3), Inches(14), Inches(1.5)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(60)
        title_para.font.bold = True
        title_para.font.color.rgb = RGBColor(255, 255, 255)
        title_para.alignment = PP_ALIGN.CENTER

        # Subtitle
        subtitle_box = slide.shapes.add_textbox(
            Inches(1), Inches(5), Inches(14), Inches(1)
        )
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = subtitle
        subtitle_para = subtitle_frame.paragraphs[0]
        subtitle_para.font.size = Pt(28)
        subtitle_para.font.color.rgb = RGBColor(226, 232, 240)
        subtitle_para.alignment = PP_ALIGN.CENTER

    def add_content_slide(title, content_items):
        """Add a content slide with bullet points"""
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # Title
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.5), Inches(15), Inches(0.8)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(40)
        title_para.font.bold = True
        title_para.font.color.rgb = PRIMARY_COLOR

        # Content
        content_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(1.8), Inches(14.4), Inches(6.2)
        )
        text_frame = content_box.text_frame
        text_frame.word_wrap = True

        for i, item in enumerate(content_items):
            if i > 0:
                text_frame.add_paragraph()
            p = text_frame.paragraphs[i]
            p.text = item
            p.font.size = Pt(20)
            p.font.color.rgb = SECONDARY_COLOR
            p.space_before = Pt(12)

    def add_two_column_slide(title, left_content, right_content):
        """Add a two-column slide"""
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # Title
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.5), Inches(15), Inches(0.8)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(40)
        title_para.font.bold = True
        title_para.font.color.rgb = PRIMARY_COLOR

        # Left column
        left_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(1.8), Inches(6.8), Inches(6.2)
        )
        left_frame = left_box.text_frame
        left_frame.word_wrap = True
        for i, item in enumerate(left_content):
            if i > 0:
                left_frame.add_paragraph()
            p = left_frame.paragraphs[i]
            p.text = item
            p.font.size = Pt(18)
            p.font.color.rgb = SECONDARY_COLOR
            p.space_before = Pt(10)

        # Right column
        right_box = slide.shapes.add_textbox(
            Inches(8.4), Inches(1.8), Inches(6.8), Inches(6.2)
        )
        right_frame = right_box.text_frame
        right_frame.word_wrap = True
        for i, item in enumerate(right_content):
            if i > 0:
                right_frame.add_paragraph()
            p = right_frame.paragraphs[i]
            p.text = item
            p.font.size = Pt(18)
            p.font.color.rgb = SECONDARY_COLOR
            p.space_before = Pt(10)

    # Slide 1: Cover
    add_title_slide("ReguAI", "Incident Response & Compliance Workflow Platform")

    # Slide 2: Problem Statement
    add_content_slide(
        "The Problem",
        [
            "🔥 Production incidents are chaotic and unstructured",
            "📊 Logs scattered across multiple systems",
            "❓ Unclear response ownership and accountability",
            "⏰ Stakeholders need constant updates",
            "📝 Post-mortems happen too late or not at all",
            "⚖️ Compliance requirements demand audit trails",
            "💼 Teams waste time on manual documentation",
        ],
    )

    # Slide 3: Solution Overview
    add_content_slide(
        "Our Solution",
        [
            "✅ Structured incident workflow from detection to resolution",
            "🤖 Automated report generation with rule-based intelligence",
            "📋 Task tracking with ownership and priorities",
            "🔍 Complete audit trail for compliance",
            "📄 Export-ready post-mortem documentation",
            "📊 Real-time dashboard with metrics and insights",
            "🔐 Secure authentication and session management",
        ],
    )

    # Slide 4: Core Features
    add_two_column_slide(
        "Core Features",
        [
            "Incident Management:",
            "• Create and track incidents",
            "• Severity classification",
            "• Status workflow tracking",
            "• Impact assessment",
            "",
            "Report Generation:",
            "• Risk classification",
            "• Timeline reconstruction",
            "• Response checklist",
            "• Technical action plans",
        ],
        [
            "Task Management:",
            "• Assign response tasks",
            "• Priority management",
            "• Status tracking",
            "• Due date monitoring",
            "",
            "Compliance & Export:",
            "• Complete audit trails",
            "• Markdown export",
            "• Stakeholder summaries",
            "• Post-mortem reports",
        ],
    )

    # Slide 5: System Architecture
    add_content_slide(
        "System Architecture",
        [
            "Frontend: Next.js 16 with App Router + React 19 + TypeScript",
            "Backend: Next.js API Routes with modular service architecture",
            "Database: PostgreSQL with Prisma ORM",
            "Authentication: Auth.js v5 (NextAuth) with JWT sessions",
            "State Management: TanStack React Query for server state",
            "UI Framework: Tailwind CSS + shadcn/ui components",
            "Validation: Zod schemas for type-safe validation",
            "Deployment: Docker Compose for local development",
        ],
    )

    # Slide 6: Technical Stack
    add_two_column_slide(
        "Technical Stack",
        [
            "Core Technologies:",
            "• Next.js 16.2.6",
            "• React 19.2.6",
            "• TypeScript 6.0.3",
            "• Tailwind CSS 4.3.0",
            "",
            "Backend:",
            "• Prisma 6.19.0",
            "• PostgreSQL",
            "• Zod 4.4.3",
            "• Auth.js 5.0.0-beta.31",
        ],
        [
            "Frontend Libraries:",
            "• TanStack React Query 5.100.10",
            "• Lucide React (icons)",
            "• shadcn/ui components",
            "",
            "Development:",
            "• ESLint 9.39.4",
            "• Docker Compose",
            "• tsx for TypeScript execution",
        ],
    )

    # Slide 7: Authentication & Security
    add_content_slide(
        "Authentication & Security",
        [
            "🔐 Auth.js v5 integration with credentials provider",
            "🎫 JWT-based session management (stateless)",
            "🛡️ Protected routes with middleware (proxy.ts)",
            "🔒 HTTP-only cookies for token storage",
            "✅ Zod validation for all API inputs",
            "🧹 Raw log sanitization (passwords, tokens, API keys)",
            "📏 Input length limits to prevent abuse",
            "📝 Comprehensive audit trail for all actions",
        ],
    )

    # Slide 8: Database Schema
    add_two_column_slide(
        "Database Models",
        [
            "User:",
            "• Identity and authentication",
            "• Role-based information",
            "",
            "Incident:",
            "• Core incident data",
            "• Severity & status tracking",
            "• Impact metrics",
            "• Timestamps & resolution",
            "",
            "IncidentAiReport:",
            "• Generated report sections",
            "• Risk classification",
            "• Action plans & checklists",
        ],
        [
            "IncidentTask:",
            "• Response work items",
            "• Owner assignment",
            "• Priority & status",
            "• Due date tracking",
            "",
            "IncidentAuditTrail:",
            "• Action logging",
            "• Actor tracking",
            "• Old/new value comparison",
            "• Timestamp records",
            "",
            "All with proper indexes",
        ],
    )

    # Slide 9: API Architecture
    add_content_slide(
        "API Design",
        [
            "RESTful API with consistent JSON responses",
            "Success: { success: true, message, data }",
            "Error: { success: false, message, errors[] }",
            "",
            "Key Endpoints:",
            "• POST /api/auth/login - Authentication",
            "• GET /api/dashboard/summary - Metrics",
            "• GET/POST/PATCH/DELETE /api/incidents - CRUD",
            "• POST /api/incidents/[id]/generate-report - Report generation",
            "• GET /api/incidents/[id]/export/markdown - Export",
        ],
    )

    # Slide 10: Report Generation Flow
    add_content_slide(
        "Automated Report Generation",
        [
            "Rule-based generation for MVP stability (no external LLM)",
            "",
            "Generated Sections:",
            "• Risk Classification - Business, technical, compliance risk",
            "• Timeline - Event reconstruction from incident data",
            "• Response Checklist - Category-based action items",
            "• Technical Action Plan - Team-specific tasks",
            "• Stakeholder Summary - Executive-friendly overview",
            "• Post-Mortem Report - Complete audit-ready documentation",
            "",
            "Auto-generation on first access - no manual trigger needed",
        ],
    )

    # Slide 11: Challenges & Solutions
    add_two_column_slide(
        "Challenges & Solutions",
        [
            "Challenge: Self-fetch issues",
            "Solution: Direct service calls in Auth.js, no HTTP self-requests",
            "",
            "Challenge: Route protection",
            "Solution: proxy.ts middleware with Auth.js wrapper",
            "",
            "Challenge: Report not generated",
            "Solution: Auto-generate on first access",
            "",
            "Challenge: Session management",
            "Solution: JWT strategy with HTTP-only cookies",
        ],
        [
            "Challenge: Type safety",
            "Solution: TypeScript + Zod validation",
            "",
            "Challenge: State management",
            "Solution: React Query for server state",
            "",
            "Challenge: Export failures",
            "Solution: Fallback export without report",
            "",
            "Challenge: Audit compliance",
            "Solution: Comprehensive audit trail",
        ],
    )

    # Slide 12: Key Improvements
    add_content_slide(
        "Key Improvements Implemented",
        [
            "✅ Auth.js v5 integration - Modern authentication",
            "✅ Proxy route protection - Secure middleware",
            "✅ Auto-report generation - Seamless UX",
            "✅ React Query migration - Optimized data fetching",
            "✅ Component refactoring - Better performance",
            "✅ Incident list optimization - Memoization & callbacks",
            "✅ Export functionality - Download reports",
            "✅ Comprehensive documentation - 15+ detailed docs",
        ],
    )

    # Slide 13: Scalability & Future
    add_content_slide(
        "Scalability & Future Roadmap",
        [
            "Current Architecture:",
            "• Stateless JWT sessions - Horizontal scaling ready",
            "• Modular service architecture - Easy to extend",
            "• Docker Compose - Consistent environments",
            "",
            "Future Enhancements:",
            "• Background job queue for report generation",
            "• WebSocket for real-time updates",
            "• AI-powered report suggestions (LLM integration)",
            "• Multi-tenant organization support",
            "• Advanced RBAC and permissions",
            "• PDF export and custom templates",
        ],
    )

    # Slide 14: Demo Flow
    add_content_slide(
        "Demo Flow",
        [
            "1. Login - Secure authentication with demo credentials",
            "2. Dashboard - View metrics and recent incidents",
            "3. Create Incident - Report a new production issue",
            "4. View Details - Automatic report generation",
            "5. Task Management - Assign and track response tasks",
            "6. Audit Trail - Complete action history",
            "7. Export - Download markdown post-mortem",
            "",
            "All features work seamlessly without manual intervention",
        ],
    )

    # Slide 15: Built with IBM Bob
    add_content_slide(
        "Built with IBM Bob",
        [
            "IBM Bob assisted with:",
            "• Architecture planning and design decisions",
            "• Database schema design and optimization",
            "• Backend module structure and organization",
            "• Validation strategy and implementation",
            "• Report generator architecture",
            "• Security review and best practices",
            "• Comprehensive documentation",
            "• Iterative refinement and debugging",
            "",
            "Result: Production-ready platform in accelerated timeline",
        ],
    )

    # Slide 16: Closing
    add_title_slide(
        "Thank You", "ReguAI - Turning Incident Chaos into Structured Compliance"
    )

    return prs


if __name__ == "__main__":
    print("Generating ReguAI presentation...")
    presentation = create_presentation()
    output_path = "project-presentation.pptx"
    presentation.save(output_path)
    print(f"✅ Presentation saved to: {output_path}")
    print(f"📊 Total slides: {len(presentation.slides)}")

# Made with Bob
